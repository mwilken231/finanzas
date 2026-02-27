#!/usr/bin/env python3
"""
Import historical transactions from the Excel file (Hoja 1).

Usage:
    # From inside the backend container or venv:
    python scripts/import_excel.py /path/to/Plan_Financiero.xlsx

    # Or via Docker:
    docker-compose exec backend python scripts/import_excel.py /data/Plan_Financiero.xlsx
"""
import sys
import os
import datetime

# Allow running from the backend/ directory
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

import openpyxl
from sqlalchemy.orm import Session

from database import engine, SessionLocal, Base
from models import Category, Transaction
from seed import run_seed


def parse_date(value):
    """Parse date from Excel cell (datetime object or ISO string)."""
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        try:
            # ISO format with timezone: 2025-08-28T00:00:00.000-03:00
            dt = datetime.datetime.fromisoformat(value.replace("Z", "+00:00"))
            return dt.date()
        except Exception:
            pass
        # Try plain date
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.datetime.strptime(value, fmt).date()
            except Exception:
                pass
    return None


def get_or_create_category(db: Session, nombre: str, tipo: str) -> Category:
    """Find existing category by name, or create a new one."""
    if not nombre:
        return None
    cat = db.query(Category).filter(Category.nombre == nombre).first()
    if cat:
        return cat
    # Create missing category
    cat = Category(nombre=nombre, tipo=tipo, orden=999)
    db.add(cat)
    db.flush()
    return cat


def import_excel(filepath: str, db: Session):
    print(f"Opening: {filepath}")
    wb = openpyxl.load_workbook(filepath, data_only=True)

    if "Hoja 1" not in wb.sheetnames:
        print("ERROR: Sheet 'Hoja 1' not found.")
        print(f"Available sheets: {wb.sheetnames}")
        sys.exit(1)

    ws = wb["Hoja 1"]
    rows = list(ws.iter_rows(values_only=True))

    if not rows:
        print("Empty sheet.")
        return

    # Skip header row
    header = rows[0]
    print(f"Header: {header}")
    data_rows = rows[1:]

    imported = 0
    skipped = 0
    errors = 0

    for i, row in enumerate(data_rows, start=2):
        if not any(v is not None for v in row):
            continue  # blank row

        fecha_n8n_raw = row[0]   # Col A: Fecha n8n
        concepto = row[1]         # Col B: Concepto
        tipo_raw = row[2]         # Col C: Tipo
        monto_raw = row[3]        # Col D: Monto
        moneda = row[4] or "ARS"  # Col E: Moneda
        categoria_raw = row[5]    # Col F: Categoría
        fecha_corr = row[6]       # Col G: fecha corregida

        # Parse tipo
        tipo = str(tipo_raw).strip().lower() if tipo_raw else None
        if tipo not in ("ingreso", "gasto"):
            skipped += 1
            continue

        # Parse monto
        try:
            monto = float(monto_raw)
            if monto <= 0:
                skipped += 1
                continue
        except (TypeError, ValueError):
            skipped += 1
            continue

        # Parse dates
        fecha = parse_date(fecha_corr) or parse_date(fecha_n8n_raw)
        if not fecha:
            print(f"  Row {i}: No valid date found, skipping. Raw: {fecha_n8n_raw!r} / {fecha_corr!r}")
            skipped += 1
            continue

        # Parse fecha_n8n as string
        fecha_n8n_str = str(fecha_n8n_raw).strip() if fecha_n8n_raw else None

        # Resolve category
        categoria_nombre = str(categoria_raw).strip() if categoria_raw else None
        categoria = None
        if categoria_nombre:
            categoria = get_or_create_category(db, categoria_nombre, tipo)

        tx = Transaction(
            fecha_n8n=fecha_n8n_str,
            fecha=fecha,
            concepto=str(concepto).strip() if concepto else None,
            tipo=tipo,
            monto=monto,
            moneda=str(moneda).strip() if moneda else "ARS",
            categoria_id=categoria.id if categoria else None,
            fuente="telegram_bot" if fecha_n8n_str else "manual",
        )
        db.add(tx)
        imported += 1

        if imported % 100 == 0:
            db.commit()
            print(f"  Imported {imported} transactions...")

    db.commit()
    print(f"\nDone! Imported: {imported} | Skipped: {skipped} | Errors: {errors}")


def main():
    if len(sys.argv) < 2:
        print("Usage: python scripts/import_excel.py <path_to_excel>")
        sys.exit(1)

    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        sys.exit(1)

    # Ensure DB is initialized
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        run_seed(db)  # Ensure categories exist
        import_excel(filepath, db)
    finally:
        db.close()


if __name__ == "__main__":
    main()
